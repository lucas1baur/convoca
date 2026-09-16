"""
Importador da planilha oficial da banca (.xlsx).

Regras de leitura, alinhadas à estrutura real da planilha do IFNMG/FADETEC:
  - Cabeçalho de dados na linha que contém "POSIÇÃO" (as linhas acima são título).
  - Chave única do candidato = INSCRIÇÃO (o CPF vem mascarado, §não-serve-de-chave).
  - Um candidato = uma linha, com várias posições preenchidas nas colunas (§7).
  - NÃO reclassifica nada: só transcreve o que a banca entregou (§45).
  - "cotista classificado em AC": a banca já esvazia as colunas de cota dele.

Nada aqui altera classificação: apenas transcreve.
"""

import openpyxl
from .db import get_conn, MODALIDADES

# Nomes de coluna esperados -> chave interna
COLMAP = {
    "POSIÇÃO": "posicao_geral",
    "INSCRIÇÃO": "inscricao",
    "CPF": "cpf",
    "NOME": "nome",
    "VAGA": "vaga",
    "MODALIDADE INSCRIÇÃO": "modalidade_inscricao",
    "NOTA": "nota",
    "SITUAÇÃO": "situacao_banca",
    "MODALIDADE CLASSIFICAÇÃO": "modalidade_classificacao",
}


def _norm(s):
    return str(s).strip().upper() if s is not None else ""


def _achar_header(ws):
    for r in range(1, 15):
        valores = [_norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "POSIÇÃO" in valores and "INSCRIÇÃO" in valores:
            return r
    raise ValueError("Não encontrei a linha de cabeçalho (POSIÇÃO / INSCRIÇÃO).")


def _to_float(v):
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _to_int(v):
    if v is None or v == "" or v == "-":
        return None
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return None


def previa(caminho_xlsx):
    """Lê a planilha e devolve um resumo, SEM gravar no banco."""
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb.active
    hr = _achar_header(ws)

    header = {}
    pos_cols = {}
    for c in range(1, ws.max_column + 1):
        nome = _norm(ws.cell(hr, c).value)
        if nome in COLMAP:
            header[COLMAP[nome]] = c
        for m in MODALIDADES:
            if nome == m.upper():
                pos_cols[m] = c

    linhas = []
    curso_nome = None
    for r in range(hr + 1, ws.max_row + 1):
        insc = ws.cell(r, header.get("inscricao", 0)).value if "inscricao" in header else None
        nome = ws.cell(r, header.get("nome", 0)).value if "nome" in header else None
        if insc is None and nome is None:
            continue
        vaga = ws.cell(r, header["vaga"]).value if "vaga" in header else None
        if vaga and curso_nome is None:
            curso_nome = str(vaga).strip()
        reg = {
            "posicao_geral": _to_int(ws.cell(r, header["posicao_geral"]).value) if "posicao_geral" in header else None,
            "inscricao": str(_to_int(insc) if isinstance(insc, float) else insc).strip() if insc is not None else None,
            "cpf": ws.cell(r, header["cpf"]).value if "cpf" in header else None,
            "nome": str(nome).strip() if nome else None,
            "vaga": str(vaga).strip() if vaga else None,
            "modalidade_inscricao": _norm(ws.cell(r, header["modalidade_inscricao"]).value) if "modalidade_inscricao" in header else None,
            "nota": _to_float(ws.cell(r, header["nota"]).value) if "nota" in header else None,
            "situacao_banca": str(ws.cell(r, header["situacao_banca"]).value).strip() if "situacao_banca" in header else None,
            "modalidade_classificacao": _norm(ws.cell(r, header["modalidade_classificacao"]).value) if "modalidade_classificacao" in header else None,
            "posicoes": {m: _to_int(ws.cell(r, col).value) for m, col in pos_cols.items()},
        }
        linhas.append(reg)

    return {
        "curso_sugerido": curso_nome,
        "total": len(linhas),
        "classificados": sum(1 for l in linhas if (l["situacao_banca"] or "").upper() == "CLASSIFICADO"),
        "lista_espera": sum(1 for l in linhas if "ESPERA" in (l["situacao_banca"] or "").upper()),
        "colunas_posicao": list(pos_cols.keys()),
        "linhas": linhas,
    }


def importar(caminho_xlsx, processo_id, curso_nome=None, curso_codigo=None):
    """Grava candidatos + elegibilidade inicial no banco. Retorna (curso_id, n)."""
    dados = previa(caminho_xlsx)
    curso_nome = curso_nome or dados["curso_sugerido"] or "Curso sem nome"

    conn = get_conn()
    c = conn.cursor()
    c.execute("""INSERT INTO curso (processo_id, codigo, nome) VALUES (?,?,?)
                 ON CONFLICT(processo_id, nome) DO UPDATE SET codigo=excluded.codigo""",
              (processo_id, curso_codigo, curso_nome))
    curso_id = c.execute("SELECT id FROM curso WHERE processo_id=? AND nome=?",
                         (processo_id, curso_nome)).fetchone()["id"]

    pos_fields = [f"pos_{m.lower()}" for m in MODALIDADES]
    n = 0
    for l in dados["linhas"]:
        if not l["inscricao"]:
            continue
        pos_vals = [l["posicoes"].get(m) for m in MODALIDADES]
        campos = ("curso_id, posicao_geral, inscricao, cpf, nome, modalidade_inscricao, "
                  "nota, situacao_banca, modalidade_classificacao, " + ", ".join(pos_fields))
        marks = ",".join("?" * (9 + len(pos_fields)))
        vals = [curso_id, l["posicao_geral"], l["inscricao"], l["cpf"], l["nome"],
                l["modalidade_inscricao"], l["nota"], l["situacao_banca"],
                l["modalidade_classificacao"]] + pos_vals
        try:
            c.execute(f"INSERT INTO candidato ({campos}) VALUES ({marks})", vals)
            cand_id = c.lastrowid
            c.execute("INSERT INTO elegibilidade (candidato_id) VALUES (?)", (cand_id,))
            n += 1
        except Exception:
            pass  # inscrição duplicada: ignora

    conn.commit()
    conn.close()
    return curso_id, n


def vagas_por_modalidade(curso_id):
    """Nº de vagas por modalidade derivado dos CLASSIFICADO (1 vaga por classificado).
    A banca já resolveu a 1ª distribuição, então cada classificado = uma vaga na
    sua modalidade_classificacao."""
    conn = get_conn()
    rows = conn.execute("""SELECT modalidade_classificacao m, COUNT(*) n
                           FROM candidato
                           WHERE curso_id=? AND UPPER(situacao_banca)='CLASSIFICADO'
                           GROUP BY modalidade_classificacao""", (curso_id,)).fetchall()
    conn.close()
    return {r["m"]: r["n"] for r in rows if r["m"]}
